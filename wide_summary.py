import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from constants import Col, FundName, SheetName, Status


YEARLY_WIDE_METRICS = ["초과건수", "초과금액", "부족건수", "부족금액", "미납건수", "미납금액"]
TOTAL_WIDE_METRICS = [
    "초과건수", "초과금액", "부족건수", "부족금액", "미납건수",
    "금액모름", "미납금액", "총건수", "총금액",
]


def _normalize_year_range(df_errors: pd.DataFrame, start_year: int | None, end_year: int | None) -> list[int]:
    years = pd.to_numeric(df_errors.get(Col.YEAR, pd.Series(dtype="object")), errors="coerce").dropna()
    if start_year is None:
        start_year = int(years.min()) if not years.empty else None
    if end_year is None:
        end_year = int(years.max()) if not years.empty else None
    if start_year is None or end_year is None or start_year > end_year:
        return []
    return list(range(int(start_year), int(end_year) + 1))


def _empty_wide_summary(group_column: str, start_year: int | None, end_year: int | None) -> pd.DataFrame:
    years = []
    if start_year is not None and end_year is not None and start_year <= end_year:
        years = list(range(int(start_year), int(end_year) + 1))
    columns = [group_column]
    columns.extend(f"{year}_{metric}" for year in years for metric in YEARLY_WIDE_METRICS)
    columns.extend(f"총계_{metric}" for metric in TOTAL_WIDE_METRICS)
    return pd.DataFrame(columns=columns)


def _add_amount_columns(df_errors: pd.DataFrame) -> pd.DataFrame:
    df = df_errors.copy()
    df[Col.YEAR] = pd.to_numeric(df[Col.YEAR], errors="coerce").astype("Int64")
    df[Col.DIFF] = pd.to_numeric(df.get(Col.DIFF), errors="coerce")
    df[Col.STANDARD] = pd.to_numeric(df.get(Col.STANDARD), errors="coerce")
    df["wide_amount"] = 0
    df["unknown_amount"] = 0

    excess_or_short = df[Col.STATUS].isin([Status.EXCESS, Status.INSUFFICIENT])
    df.loc[excess_or_short, "wide_amount"] = df.loc[excess_or_short, Col.DIFF].abs().fillna(0)

    unpaid = df[Col.STATUS] == Status.UNPAID
    standard_known = df[Col.STANDARD].notna()
    df.loc[unpaid & standard_known, "wide_amount"] = df.loc[unpaid & standard_known, Col.STANDARD].fillna(0)
    df.loc[unpaid & ~standard_known, "unknown_amount"] = 1
    return df


def _build_wide_summary(
    df_errors: pd.DataFrame,
    group_column: str,
    start_year: int | None = None,
    end_year: int | None = None,
) -> pd.DataFrame:
    if df_errors is None or df_errors.empty:
        return _empty_wide_summary(group_column, start_year, end_year)

    years = _normalize_year_range(df_errors, start_year, end_year)
    columns = [group_column]
    columns.extend(f"{year}_{metric}" for year in years for metric in YEARLY_WIDE_METRICS)
    columns.extend(f"총계_{metric}" for metric in TOTAL_WIDE_METRICS)

    df = _add_amount_columns(df_errors)
    rows = []
    for group_value, group_df in df.groupby(group_column, dropna=False):
        ranged_group_df = group_df[group_df[Col.YEAR].isin(years)]
        row = {column: 0 for column in columns}
        row[group_column] = group_value
        for year in years:
            year_df = ranged_group_df[ranged_group_df[Col.YEAR] == year]
            for status, count_metric, amount_metric in [
                (Status.EXCESS, "초과건수", "초과금액"),
                (Status.INSUFFICIENT, "부족건수", "부족금액"),
                (Status.UNPAID, "미납건수", "미납금액"),
            ]:
                status_df = year_df[year_df[Col.STATUS] == status]
                row[f"{year}_{count_metric}"] = int(len(status_df))
                row[f"{year}_{amount_metric}"] = int(status_df["wide_amount"].sum())

        for status, count_metric, amount_metric in [
            (Status.EXCESS, "초과건수", "초과금액"),
            (Status.INSUFFICIENT, "부족건수", "부족금액"),
            (Status.UNPAID, "미납건수", "미납금액"),
        ]:
            status_df = ranged_group_df[ranged_group_df[Col.STATUS] == status]
            row[f"총계_{count_metric}"] = int(len(status_df))
            row[f"총계_{amount_metric}"] = int(status_df["wide_amount"].sum())
        row["총계_금액모름"] = int(ranged_group_df["unknown_amount"].sum())
        row["총계_총건수"] = int(row["총계_초과건수"] + row["총계_부족건수"] + row["총계_미납건수"])
        row["총계_총금액"] = int(row["총계_초과금액"] + row["총계_부족금액"] + row["총계_미납금액"])
        rows.append(row)

    summary = pd.DataFrame(rows, columns=columns)
    if group_column == Col.NAME:
        return summary.sort_values(["총계_총건수", "총계_총금액", Col.NAME], ascending=[False, False, True]).reset_index(drop=True)

    fund_order = {FundName.OPERATING: 0, FundName.COOPERATION: 1, FundName.WELFARE: 2}
    summary["_fund_order"] = summary[group_column].map(fund_order).fillna(99)
    return summary.sort_values(["_fund_order", group_column]).drop(columns=["_fund_order"]).reset_index(drop=True)


def build_wide_error_summary_by_person(
    df_errors: pd.DataFrame,
    start_year: int | None = None,
    end_year: int | None = None,
) -> pd.DataFrame:
    return _build_wide_summary(df_errors, Col.NAME, start_year, end_year)


def build_wide_error_summary_by_fund(
    df_errors: pd.DataFrame,
    start_year: int | None = None,
    end_year: int | None = None,
) -> pd.DataFrame:
    return _build_wide_summary(df_errors, Col.FUND_NAME, start_year, end_year)


def write_wide_summary_sheets(writer, df_errors: pd.DataFrame, start_year=None, end_year=None) -> None:
    person_status = build_wide_error_summary_by_person(df_errors, start_year, end_year)
    fund_status = build_wide_error_summary_by_fund(df_errors, start_year, end_year)
    _write_wide_summary_sheet(writer, SheetName.PERSON_ERROR_STATUS, Col.NAME, person_status)
    _write_wide_summary_sheet(writer, SheetName.FUND_ERROR_STATUS, Col.FUND_NAME, fund_status)


def _write_wide_summary_sheet(writer, sheet_name: str, group_column: str, summary_df: pd.DataFrame) -> None:
    if summary_df is None:
        return

    ws = writer.book.create_sheet(sheet_name)
    years = _extract_years(summary_df)
    layout = _build_sheet_layout(group_column, years)
    _write_headers(ws, group_column, years)
    _write_data_rows(ws, summary_df, layout, group_column)
    _style_sheet(ws, layout)


def _extract_years(summary_df: pd.DataFrame) -> list[int]:
    years = []
    for column in summary_df.columns:
        text = str(column)
        if "_" not in text:
            continue
        prefix = text.split("_", 1)[0]
        if prefix.isdigit() and int(prefix) not in years:
            years.append(int(prefix))
    return years


def _build_sheet_layout(group_column: str, years: list[int]) -> list[tuple[str, int | str | None, str | None]]:
    layout: list[tuple[str, int | str | None, str | None]] = [("field", group_column, group_column), ("sep", None, None)]
    for year in years:
        for metric in YEARLY_WIDE_METRICS:
            layout.append(("year", year, metric))
        layout.append(("sep", None, None))
    for metric in TOTAL_WIDE_METRICS:
        layout.append(("total", "총계", metric))
    return layout


def _write_headers(ws, group_column: str, years: list[int]) -> None:
    ws.cell(row=1, column=1, value=group_column)
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    col_idx = 3
    for year in years:
        start_col = col_idx
        end_col = col_idx + len(YEARLY_WIDE_METRICS) - 1
        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        ws.cell(row=1, column=start_col, value=year)
        for offset, metric in enumerate(YEARLY_WIDE_METRICS):
            ws.cell(row=2, column=start_col + offset, value=metric)
        col_idx = end_col + 2

    total_start = col_idx
    total_end = total_start + len(TOTAL_WIDE_METRICS) - 1
    ws.merge_cells(start_row=1, start_column=total_start, end_row=1, end_column=total_end)
    ws.cell(row=1, column=total_start, value="총계")
    for offset, metric in enumerate(TOTAL_WIDE_METRICS):
        ws.cell(row=2, column=total_start + offset, value=metric)


def _write_data_rows(ws, summary_df: pd.DataFrame, layout, group_column: str) -> None:
    for row_offset, (_, source_row) in enumerate(summary_df.iterrows(), start=3):
        for column, item in enumerate(layout, start=1):
            kind, group, metric = item
            if kind == "field":
                value = source_row[group_column]
            elif kind == "sep":
                value = None
            elif kind == "year":
                value = source_row.get(f"{group}_{metric}", 0)
            else:
                value = source_row.get(f"총계_{metric}", 0)
            ws.cell(row=row_offset, column=column, value=_escape_excel_formula(value))


def _escape_excel_formula(value):
    if not isinstance(value, str):
        return value
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{value}"
    return value


def _style_sheet(ws, layout) -> None:
    fills = {
        "normal": PatternFill(start_color="FF272F3A", end_color="FF272F3A", fill_type="solid"),
        "total": PatternFill(start_color="FF963634", end_color="FF963634", fill_type="solid"),
        "separator": PatternFill(start_color="FFBFBFBF", end_color="FFBFBFBF", fill_type="solid"),
        "even": PatternFill(start_color="FFF5F3F4", end_color="FFF5F3F4", fill_type="solid"),
        "total_odd": PatternFill(start_color="FFFAF0F0", end_color="FFFAF0F0", fill_type="solid"),
        "total_even": PatternFill(start_color="FFF5E3E3", end_color="FFF5E3E3", fill_type="solid"),
    }
    header_row_1_font = Font(color="FFFFFFFF", bold=True, size=10)
    header_row_2_font = Font(color="FFFFFFFF", bold=True, size=9)
    name_font = Font(color="FF000000", bold=False, size=11)
    metric_fonts = {
        "초과건수": Font(color="FF059669", bold=True, size=11),
        "초과금액": Font(color="FF059669", bold=True, size=11),
        "부족건수": Font(color="FFEF4444", bold=True, size=11),
        "부족금액": Font(color="FFEF4444", bold=True, size=11),
        "미납건수": Font(color="FF595959", bold=True, size=11),
        "금액모름": Font(color="FF595959", bold=True, size=11),
        "미납금액": Font(color="FF595959", bold=True, size=11),
        "총건수": Font(color="FF963634", bold=True, size=11),
        "총금액": Font(color="FF963634", bold=True, size=11),
    }
    thin = Side(style="thin", color="FF9CA3AF")
    thick = Side(style="thick", color="FF595959")
    amount_metrics = {"초과금액", "부족금액", "미납금액", "총금액"}
    thick_left_metrics = {"초과건수", "부족건수", "미납건수", "총건수"}
    thick_right_metrics = {"초과금액", "부족금액", "미납금액", "총금액"}
    max_col = len(layout)

    for column, item in enumerate(layout, start=1):
        kind, _, metric = item
        letter = get_column_letter(column)
        if kind == "sep":
            ws.column_dimensions[letter].width = 3.0
        elif kind == "year" and metric in amount_metrics:
            ws.column_dimensions[letter].width = 9.72
        elif kind == "total" and metric in amount_metrics:
            ws.column_dimensions[letter].width = 10.8
        elif kind == "year" or kind == "total":
            ws.column_dimensions[letter].width = 8.64
        else:
            ws.column_dimensions[letter].width = 9.6

        for row in range(1, ws.max_row + 1):
            cell = ws.cell(row=row, column=column)
            left = thick if column == 1 or kind == "sep" or metric in thick_left_metrics else thin
            right = thick if kind == "sep" or metric in thick_right_metrics or column == max_col else thin
            top = thick if row == 1 else thin
            bottom = thick if row == ws.max_row else thin
            cell.border = Border(left=left, right=right, top=top, bottom=bottom)

            if row <= 2:
                cell.font = header_row_1_font if row == 1 else header_row_2_font
                cell.fill = fills["total"] if kind == "total" else fills["normal"]
                if kind == "sep":
                    cell.fill = fills["separator"]
                cell.alignment = Alignment(
                    horizontal="center" if kind == "field" or row == 1 else "left",
                    vertical="center",
                )
                continue

            _style_data_cell(cell, kind, metric, row, name_font, metric_fonts, fills, amount_metrics)

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 24
    ws.auto_filter.ref = f"A2:{get_column_letter(max_col)}{ws.max_row}"
    ws.freeze_panes = "B3"


def _style_data_cell(cell, kind, metric, row, name_font, metric_fonts, fills, amount_metrics) -> None:
    is_even_data_row = (row - 2) % 2 == 0
    if kind == "sep":
        cell.fill = fills["separator"]
        return
    if metric in {"총건수", "총금액"}:
        cell.fill = fills["total_even"] if is_even_data_row else fills["total_odd"]
    elif is_even_data_row:
        cell.fill = fills["even"]
    cell.font = name_font if kind == "field" else metric_fonts.get(metric, name_font)
    cell.alignment = Alignment(horizontal="right" if metric in amount_metrics else "center", vertical="center")
    if isinstance(cell.value, (int, float)):
        cell.number_format = "#,##0"
