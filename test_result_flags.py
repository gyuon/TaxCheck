import unittest
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import data_processor
from constants import Col, RESULT_COLUMNS, SheetName, Status


class TestResultFlags(unittest.TestCase):
    def test_error_target_marks_unpaid_and_thousand_won_boundaries(self):
        rows = []
        for name, cooperation_amount in {
            "PositiveBoundary": 41000,
            "PositiveInside": 40999,
            "NegativeBoundary": 39000,
            "NegativeInside": 39001,
        }.items():
            rows.extend([
                {
                    Col.NAME: name,
                    Col.YEAR: 2021,
                    Col.MONTH: 1,
                    Col.CODE1: 1,
                    Col.CODE2: 1,
                    Col.RAW_DEPOSIT: 100000,
                },
                {
                    Col.NAME: name,
                    Col.YEAR: 2021,
                    Col.MONTH: 1,
                    Col.CODE1: 2,
                    Col.CODE2: 1,
                    Col.RAW_DEPOSIT: cooperation_amount,
                },
            ])

        errors = data_processor.detect_errors(pd.DataFrame(rows)).set_index(Col.NAME)

        self.assertEqual(errors.loc["PositiveBoundary", Col.DIFF], 1000)
        self.assertEqual(errors.loc["PositiveBoundary", Col.ERROR_TARGET], "O")
        self.assertEqual(errors.loc["NegativeBoundary", Col.DIFF], -1000)
        self.assertEqual(errors.loc["NegativeBoundary", Col.ERROR_TARGET], "O")
        self.assertEqual(errors.loc["PositiveInside", Col.ERROR_TARGET], "")
        self.assertEqual(errors.loc["NegativeInside", Col.ERROR_TARGET], "")

        unpaid_source = pd.DataFrame([
            {
                Col.NAME: "Unpaid",
                Col.YEAR: 2021,
                Col.MONTH: 1,
                Col.CODE1: 1,
                Col.CODE2: 1,
                Col.RAW_DEPOSIT: 100000,
            }
        ])
        missed, _ = data_processor.generate_missed_months(unpaid_source)
        self.assertTrue((missed[Col.ERROR_TARGET] == "O").all())

    def test_three_month_unpaid_flag_keeps_existing_boundary_rule(self):
        cases = [
            (2026, 1, "※정리본_26.4.14.xlsx", "O"),
            (2026, 2, "※정리본_26.4.14.xlsx", ""),
            (2025, 10, "※정리본_26.1.14.xlsx", "O"),
        ]

        for year, month, filename, expected in cases:
            with self.subTest(year=year, month=month, filename=filename):
                source = pd.DataFrame([
                    {
                        Col.NAME: "Unpaid",
                        Col.YEAR: year,
                        Col.MONTH: month,
                        Col.CODE1: 1,
                        Col.CODE2: 1,
                        Col.RAW_DEPOSIT: 100000,
                    }
                ])
                missed, _ = data_processor.generate_missed_months(source, filename=filename)

                self.assertFalse(missed.empty)
                self.assertTrue((missed[Col.UNPAID_3_MONTHS] == expected).all())

    def test_error_result_sheet_has_new_columns_and_inactive_filter(self):
        df_first = pd.DataFrame({Col.NAME: ["A"], Col.YEAR: [2021], Col.MONTH: [1]})
        source = pd.DataFrame([
            {
                Col.NAME: "A",
                Col.YEAR: 2021,
                Col.MONTH: 1,
                Col.CODE1: 1,
                Col.CODE2: 1,
                Col.RAW_DEPOSIT: 100000,
            },
            {
                Col.NAME: "A",
                Col.YEAR: 2021,
                Col.MONTH: 1,
                Col.CODE1: 2,
                Col.CODE2: 1,
                Col.RAW_DEPOSIT: 41000,
            },
        ])
        errors = data_processor.detect_errors(source)

        excel_bytes, _ = data_processor.to_excel_bytes(
            df_first,
            errors,
            "result.xlsx",
            df_raw=pd.DataFrame(),
        )
        workbook = load_workbook(BytesIO(excel_bytes))
        worksheet = workbook[SheetName.ERROR_RESULT]
        headers = [cell.value for cell in worksheet[1]]

        self.assertEqual(headers, ["번호", *RESULT_COLUMNS])
        self.assertIn(Col.ERROR_TARGET, headers)
        self.assertIn(Col.UNPAID_3_MONTHS, headers)
        self.assertNotIn("비고", headers)
        self.assertEqual(
            worksheet.auto_filter.ref,
            f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}",
        )
        self.assertEqual(worksheet.auto_filter.filterColumn, [])

        error_target_column = headers.index(Col.ERROR_TARGET) + 1
        long_unpaid_column = headers.index(Col.UNPAID_3_MONTHS) + 1
        self.assertEqual(worksheet.cell(2, error_target_column).value, "O")
        self.assertEqual(worksheet.cell(2, error_target_column).alignment.horizontal, "center")
        self.assertEqual(worksheet.cell(2, long_unpaid_column).alignment.horizontal, "center")

    def test_excel_export_migrates_legacy_result_columns(self):
        legacy_errors = pd.DataFrame([
            {
                Col.NAME: "Legacy",
                Col.YEAR: 2021,
                Col.MONTH: 1,
                Col.FUND_NAME: "협력",
                Col.CODE: "21",
                Col.STATUS: Status.EXCESS,
                Col.DIFF: 1000,
                Col.DEPOSIT: 41000,
                Col.STANDARD: 40000,
                Col.FORMULA: "",
                Col.PREV_MONTH_AMT: 0,
                Col.NEXT_MONTH_AMT: 0,
                "비고": "3개월 이전 미납분",
            }
        ])

        excel_bytes, _ = data_processor.to_excel_bytes(
            pd.DataFrame({Col.NAME: ["Legacy"], Col.YEAR: [2021], Col.MONTH: [1]}),
            legacy_errors,
            "legacy.xlsx",
            df_raw=pd.DataFrame(),
        )
        worksheet = load_workbook(BytesIO(excel_bytes))[SheetName.ERROR_RESULT]
        headers = [cell.value for cell in worksheet[1]]

        self.assertEqual(headers, ["번호", *RESULT_COLUMNS])
        self.assertEqual(worksheet.cell(2, headers.index(Col.ERROR_TARGET) + 1).value, "O")
        self.assertEqual(worksheet.cell(2, headers.index(Col.UNPAID_3_MONTHS) + 1).value, "O")


if __name__ == "__main__":
    unittest.main()
